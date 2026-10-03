import json
from openpyxl import load_workbook

SOURCE = "/workspace/scratch/b3a85e99bc55/upload/BÁO GIÁ  TỰ ĐỘNG 2026_03092026.xlsm"
OUTPUT = "/workspace/sites/bach-ngan-quote/dist/products.json"

wb = load_workbook(SOURCE, read_only=True, data_only=True)

packing = {}
ws_pack = wb["QUY CÁCH THÙNG"]
for row in ws_pack.iter_rows(min_row=6, max_col=22, values_only=True):
    sku = str(row[1]).strip() if row[1] is not None else ""
    if not sku:
        continue
    per_carton = row[16] if isinstance(row[16], (int, float)) else 0
    carton_weight = row[21] if isinstance(row[21], (int, float)) else 0
    packing[sku] = [per_carton, carton_weight]

products = {}
ws = wb["GIÁ L1"]
for row in ws.iter_rows(min_row=2, max_col=11, values_only=True):
    sku = str(row[0]).strip() if row[0] is not None else ""
    name = str(row[1]).strip() if row[1] is not None else ""
    if not sku or not name:
        continue
    price = next((v for v in (row[10], row[9], row[7]) if isinstance(v, (int, float))), 0)
    retail = row[7] if isinstance(row[7], (int, float)) else 0
    pack = packing.get(sku, [0, 0])
    products[sku] = {
        "s": sku,
        "n": name,
        "b": str(row[2]).strip() if row[2] is not None else "",
        "p": str(row[3]).strip() if row[3] is not None else "",
        "u": str(row[4]).strip() if row[4] is not None else "",
        "t": str(row[5]).strip() if row[5] is not None else "",
        "r": retail,
        "g": price,
        "c": pack[0],
        "w": pack[1],
    }

with open(OUTPUT, "w", encoding="utf-8") as handle:
    json.dump(list(products.values()), handle, ensure_ascii=False, separators=(",", ":"))

print(f"Exported {len(products)} products to {OUTPUT}")
